"""Correct RAVDESS S5 response semantics, anchored by primary descriptors/accuracy."""
import csv
import hashlib
import json
from pathlib import Path
import openpyxl

HERE = Path(__file__).resolve().parent
WORKBOOK = HERE.parent / 'recovery_2026_10_10/external_data/ravdess/pone.0196391.s008.xlsx'
LABELS = ('angry','disgust','fear','happy','neutral','sad','surprise','calm','none')
RESPONSE_IDS = (5,7,6,3,1,4,8,2,9)
INTENDED = {1:'Neutral',2:'Calm',3:'Happy',4:'Sad',5:'Angry',6:'Fearful',7:'Disgust',8:'Surprise'}


def main():
    book = openpyxl.load_workbook(WORKBOOK,read_only=True,data_only=True)
    accuracy = {r[0]:r for r in book['Speech_RawAccuracy'].iter_rows(min_row=4,values_only=True)}
    records=[];verified=0;positive_semantics={i:0 for i in range(1,9)}
    for row in book['Speech_EmoCategory'].iter_rows(min_row=4,values_only=True):
        identifiers=[int(v) for v in Path(row[0]).stem.split('-')]
        assert identifiers==[int(v) for v in row[8:15]]
        intended=int(row[10])
        assert row[3]==INTENDED[intended],(row[0],row[3],intended)
        assert accuracy[row[0]][:15]==row[:15]
        responses=tuple(int(v) for v in row[15:25])
        assert all(1<=v<=9 for v in responses)
        expected=tuple(int(v==intended or (intended in (1,2) and v in (1,2))) for v in responses)
        assert expected==accuracy[row[0]][15:25]
        positive_semantics[intended]+=sum(v==intended and correct==1 for v,correct in zip(responses,expected))
        verified+=1
        if int(row[8])!=2:continue
        assert row[1]=='VO' and int(row[9])==1
        counts=[responses.count(identifier) for identifier in RESPONSE_IDS]
        assert sum(counts)==10
        records.append(dict(filename=row[0],actor=int(row[14]),annotation_n=10,
                            **{'count_'+label:value for label,value in zip(LABELS,counts)},
                            **{'q_'+label:value/10 for label,value in zip(LABELS,counts)}))
    book.close()
    assert verified==4320 and len(records)==1440 and len({r['filename'] for r in records})==1440
    assert all(value>0 for value in positive_semantics.values())
    records.sort(key=lambda r:r['filename'])
    target=HERE/'video_only_speech_votes.csv'
    with target.open('w',newline='',encoding='utf-8') as stream:
        writer=csv.DictWriter(stream,fieldnames=list(records[0]));writer.writeheader();writer.writerows(records)
    report=dict(source_workbook_sha256=hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(),
                output_sha256=hashlib.sha256(target.read_bytes()).hexdigest(),
                categorical_accuracy_rows_verified=verified,video_only_clips=1440,
                response_ids=dict(zip(LABELS,RESPONSE_IDS)),intended_descriptor_ids=INTENDED,
                positive_response_accuracy_anchors=positive_semantics,
                correction='Prior metadata adapter transposed Disgust/Surprise. Corrected before scientific external outcomes. Raw workbook and prior edition preserved.',
                source='https://doi.org/10.1371/journal.pone.0196391',
                official_filename_legend='https://zenodo.org/records/1188976')
    (HERE/'annotation_mapping_audit.json').write_text(json.dumps(report,indent=2))
    print(json.dumps(report,indent=2))


if __name__=='__main__':main()
